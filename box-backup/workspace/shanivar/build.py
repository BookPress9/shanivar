from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import csv, openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

d=Document()
st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5)
def H(t,l=1): d.add_heading(t,l)
def P(t,b=False,i=False):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=b; r.italic=i; return p
def B(items):
    for x in items: d.add_paragraph(x,style='List Bullet')
def N(items):
    for x in items: d.add_paragraph(x,style='List Number')
def T(head,rows):
    t=d.add_table(rows=1,cols=len(head)); t.style='Light Grid Accent 1'
    for c,h in zip(t.rows[0].cells,head):
        c.text=''; c.paragraphs[0].add_run(h).bold=True
    for r in rows:
        cs=t.add_row().cells
        for c,v in zip(cs,r): c.text=v
    d.add_paragraph()
def BOX(t):
    p=P(t,b=True); p.runs[0].font.color.rgb=RGBColor(0xB0,0x30,0x00)

d.add_heading('Shanivar Shikhar — Launch Package',0)
P('Trek + sit morning · Jaipur · Sun 18 Oct 2026 · ~6:15–8:30 IST · Cap 8 · Paid test ₹399 · Host: Shubham Kumar · with [PARTNER]')
BOX('Collaboration: Shanivar Shikhar × [PARTNER]')
P('Roles/split (TO CONFIRM): Shubham — route, hosting, safety, UPI collection; [PARTNER] — promotion + helper on the day; revenue/cost split [X% / Y%].')
BOX('GO/NO-GO GATE — Fri 16 Oct, 6 pm IST')
B(['Written YES from DCF Wildlife Jaipur (Jal Mahal ridge, Nahargarh) → Option A at ₹399.',
   'No written yes → move to Option B (non-sanctuary, recce-confirmed) at same ₹399, no refund needed — or postpone with full refund.',
   'Never run the paid group on sanctuary land without written permission. Never say or imply permission exists.',
   'Tell buyers: “Route: Jal Mahal ridge or Sumel/Jamdoli side — final by Fri 16 Oct.”'])
T(['','Option A — Jal Mahal ridge','Option B — Sumel/Jamdoli fallback'],[
 ['Access','Park Kanak Vrindavan side; ~1.5 km to ridge; skip temple stair','Past Sisodia Rani ka Bagh; distance unclear (blogs: 1 km vs 3–4 km)'],
 ['Land','Nahargarh hills / sanctuary — paid group needs written permission','Treated as non-sanctuary — confirm on recce'],
 ['Permission','DCF Wildlife Jaipur, Jagdish Chand Gupta · 9413807116 · dcfwl.jpr.forest@rajasthan.gov.in','Confirm no permission needed'],
 ['Status','NOT confirmed','Recce needed'],
 ['Source','thehumbleworld.com/2022/09/jal-mahal-trek-jaipur.html','Blogs (conflicting)']])
P('Note: no public day-hike permit exists for Nahargarh; unlicensed weekend groups reported (First India, 14 Apr 2025).',i=True)

H('1. Instagram')
H('Captions (pick one)',2)
P('A —',b=True)
P('Walk up quietly. Sit 10 minutes. Walk down.\nNo music. No mantra. No belief needed.\nSun 18 Oct · 6:15 am · Jaipur · [MEETING POINT – TO CONFIRM]\n₹399 · water + light refreshments incl. · 8 seats\nWith [PARTNER]. DM “SHIKHAR” to join.')
P('B —',b=True)
P('Out of your head. Into your legs.\nA short hill walk at sunrise, a quiet sit at the top, then back down.\nSun 18 Oct · 6:15 am · [MEETING POINT – TO CONFIRM] · ₹399 incl. water + refreshments · 8 seats · DM or [LINK]\nWith [PARTNER].')
P('C —',b=True)
P('Move. Feel. Notice. That’s it.\nShanivar Shikhar — one sunrise morning, 8 people, ~2 hrs.\nSun 18 Oct · 6:15 am · [MEETING POINT – TO CONFIRM] · ₹399 · Register: [LINK] / DM\nShanivar Shikhar × [PARTNER]')
P('Add to all: “Transport not included. Route final by Fri 16 Oct.”',i=True)
P('Co-branding line: “Shanivar Shikhar × [PARTNER] — a quiet sunrise walk, made together.” Credit: “with [PARTNER]”; tag @[PARTNER] in post + stories.',i=True)
H('Stories (3 frames)',2)
T(['Frame','Text'],[
 ['1','Sunday. Sunrise. A hill in Jaipur.'],
 ['2','Walk up quietly → sit 10 min at the top → walk down.\nNo music. No mantra. No sharing circle.'],
 ['3','Sun 18 Oct · 6:15 am · ₹399 · 8 seats\nWater + light refreshments incl.\n[MEETING POINT – TO CONFIRM]\nwith [PARTNER] · DM “SHIKHAR” / link sticker']])
H('Hashtags',2)
P('#Shanivar #ShanivarShikhar #[PARTNER] #JaipurMornings #JaipurTrek #SunriseJaipur #SlowHike #WalkingMeditation #JaipurOutdoors #AravalliHills #PhonesDown')
H('Visual direction',2)
B(['No real people. No faces, no silhouettes of identifiable participants.',
   'Shots: empty trail at first light; rock ledge at the top; shoes on dust (stock/illustrated only); wide ridge + soft sunrise haze; water bottle on a rock.',
   'Palette: dusty ochre, stone grey, pale morning blue. Plain sans-serif text. No om symbols, no lotus, no glow effects.',
   'Do not show the fort or restricted areas; do not geotag the ridge until route is confirmed.'])

H('2. WhatsApp broadcast')
P('Hi! Small thing I’m starting with [PARTNER]: Shanivar Shikhar.\nSun 18 Oct, 6:15–8:30 am, Jaipur.\nEasy hill walk up, 10 min quiet sit at the top, walk down. No music, no mantra, nothing to believe.\nMax 8 people. Meeting point: [TO CONFIRM].\n₹399 — water + light refreshments incl. (transport not incl.)\nReply “IN” for the form; seat confirmed after UPI.')

H('3. Itinerary (approx. — sunrise Jaipur mid-Oct ~6:25 IST, approximate)')
T(['Time (IST)','What','Notes'],[
 ['6:00','Host + helper at meeting point','Head-count list, water (12 L), refreshments, first-aid'],
 ['6:10–6:20','Arrivals','Tick names. Hand each person 1 L water. Last call 6:20.'],
 ['6:20–6:25','Welcome + safety brief','Opening line; route; stay together; monkeys/leopard rules'],
 ['6:25–6:32','Optional light mobility warm-up','Ankles, hips, shoulders. Opt-in.'],
 ['6:32','Head-count #1 · start walk','Host front, helper back'],
 ['6:32–7:10','Walk up quietly','Steady pace; one water stop ~6:50'],
 ['7:10–7:15','Arrive top · settle','Find a spot; phones away'],
 ['7:15–7:27','Seated sit (10–12 min)','Eyes open or closed. No mantra, no guidance beyond start/end.'],
 ['7:27–7:35','Stretch, water, quiet look around','Phones back allowed. No sharing circle.'],
 ['7:35','Head-count #2 · start descent',''],
 ['7:35–8:15','Walk down','Slower on loose rock'],
 ['8:15','Head-count #3 at base',''],
 ['8:15–8:25','Light refreshments at base','Fruit, dates/chana, nimbu pani. Quiet, no music.'],
 ['8:25–8:30','Close','Closing line; feedback link; next date. Done by 8:30.']])
P('Option B: adjust walk times after recce if distance is 3–4 km.',i=True)

H('4. Bring / wear / safety')
T(['Bring','Wear','Don’t bring'],[
 ['Water: 1 L given at start (own bottle optional)','Closed shoes with grip','Speakers / music'],
 ['Light refreshments provided at the end','Light, breathable layers','Food to feed animals'],
 ['Cap, sunscreen','Something to sit on (optional)','Valuables, open snacks in hand'],
 ['Any personal meds','No flip-flops','Pets, alcohol']])
H('Safety & expectations',2)
B(['Easy–moderate walk, rocky in parts. You walk at your own risk; tell us about any medical condition.',
   'Stay with the group. Host front, helper back. No solo detours.',
   'Wildlife: leopards live in these hills, monkeys are common. Don’t feed, don’t stare, keep food packed, stay together.',
   'Heat: we finish by 8:30 to avoid it. Drink water.',
   'No night start. No entering the fort or any restricted/fenced area.',
   'Phones away only during the sit. Otherwise your choice.',
   'This is a walk and a quiet sit. Not a class, ritual or workshop. No belief needed.',
   'No insurance is provided.'])
H('Cancellation / weather',2)
P('If rain, storm or heat alert is forecast, or route access is unclear, host cancels by 8 pm Sat 17 Oct on WhatsApp → full ₹399 refund by UPI within 48 h. Participant cancels before Fri 16 Oct 6 pm → full refund; after that → no refund (seat transferable). Route change A→B → no refund.')

H('5. Registration flow')
H('Form fields',2)
N(['Full name','WhatsApp number','Age (18+ only)','Fitness level: easy / moderate / regular hiker',
   'Medical conditions we should know (or “none”)','Emergency contact (name + number)','How did you hear about us?',
   'Tick: I understand this is a walk at my own risk, no insurance provided.',
   'Tick: I’ll pay ₹399 by UPI; seat confirmed only after payment.'])
H('Payment (UPI)',2)
N(['Form received → send: “Pay ₹399 to UPI [UPI-ID] (name: [NAME]). Send screenshot or UTR/ref no. here.”',
   'Check UPI app; log UTR/ref + screenshot Y/N in tracker (Paid = 399).',
   'Only then send confirmation. Unpaid after 12 h → seat goes to waitlist.'])
H('Messages',2)
T(['When','Message'],[
 ['Confirmation','Hi [Name], ₹399 received (ref [UTR]) — you’re in for Shanivar Shikhar × [PARTNER], Sun 18 Oct, 6:15 am at [MEETING POINT]. Water + light refreshments provided. Wear closed shoes, cap. Route final by Fri 16 Oct. Reply “OK”. — Shubham'],
 ['24h before (Sat 6:15 am)','Tomorrow, 6:15 am sharp at [MEETING POINT] [map link]. Route: [A/B]. Water given; grippy shoes, cap. Done by 8:30. Can’t come? Tell me today so I can give your seat away.'],
 ['2h before (Sun 4:15 am — schedule)','Morning! See you 6:15 at [MEETING POINT]. We start walking 6:32. Running late? Message me before 6:15.'],
 ['Waitlist','Hi [Name], we’re full (8). You’re #[N] on the waitlist. I’ll message if a seat opens — pay ₹399 by UPI within 2 h to take it. Latest Sat 8 pm.'],
 ['Thank-you + feedback','Thank you for walking with us today. 3 quick questions:\n1. What did you notice most?\n2. What would you change?\n3. Would you come again? (yes / maybe / no)\nNext one: [Sat/Mon DATE], same time. Reply “IN” to hold a seat. — Shubham']])
H('Rules',2)
B(['Waitlist: first come, first served. Seat offered by WhatsApp; 2 hrs to accept + pay, then next person. Closes Sat 17 Oct, 8 pm.',
   'No-show: wait till 6:20, then leave. No refund; no priority next time.',
   'Cap 8; minimum 6 paid by Fri 16 Oct 6 pm, else postpone + full refund.'])

H('6. Host run sheet')
H('Before',2)
B(['☐ Fri 16 Oct 6 pm: GO/NO-GO decided (written permission or fallback/postpone)',
   '☐ Option B: recce done (distance, time up/down, sit spot, parking)',
   '☐ Meeting point + map pin fixed and sent',
   '☐ Helper confirmed (back of group) · [PARTNER] roles/split confirmed in writing',
   '☐ All seats paid; UPI refs match tracker',
   '☐ Buy water: 8 × 1 L bottles + 4 L spare (12 L). Host carries 2 L spare + refreshments; helper carries 2 L spare.',
   '☐ Buy refreshments: fruit (bananas/oranges), dates or roasted chana, nimbu pani in a flask, paper cups, trash bag',
   '☐ Cost (ESTIMATE, not verified): water ~₹250 + refreshments ~₹600–900 ≈ ₹850–1,150 for 8 vs revenue ₹3,192',
   '☐ Weather/heat check Sat 8 pm',
   '☐ Tracker updated; reminders scheduled',
   '☐ Pack: first-aid kit, 2 L spare water, phone charged, power bank, whistle, printed list + emergency contacts, small timer/bell'])
H('During',2)
B(['☐ 6:00 on site · ☐ head-count on arrival (tick in tracker)',
   '☐ Opening: “Welcome. We’ll walk up quietly, sit for about ten minutes at the top, and walk down. Nothing to do or believe — just notice.”',
   '☐ Hand out 1 L water each · Safety brief (stay together, wildlife, water, no fort/restricted areas)',
   '☐ Keep refreshments packed till base (monkeys)',
   '☐ Head-count #1 before walk · #2 before descent · #3 at base → serve refreshments',
   '☐ Start/end the sit with a word or soft bell. No talk-over.',
   '☐ Closing: “That’s it. Thank you for coming out this morning. Take your time getting back.”'])
H('After',2)
B(['☐ Everyone left safely (final count matches) · all cups/wrappers packed out',
   '☐ Log actual spend; settle split with [PARTNER]',
   '☐ Thank-you + feedback sent same day by 12 pm',
   '☐ Tracker: arrived, paid, feedback columns updated',
   '☐ 3 notes for next time; leave no trace check done'])

H('7. Tracker template')
T(['Name','Number','Status','Paid (₹399)','UPI ref / screenshot','Arrived','Feedback'],
  [['','','Confirmed / Waitlist / Cancelled','399 / 0 / Refunded','UTR no. / Y','Y / N',''] ]+[['']*7 for _ in range(7)])
P('Live file: shanivar_tracker_18oct.xlsx (also .csv).',i=True)

H('8. Open items for Shubham')
T(['Item','Status','Decide by'],[
 ['Meeting point','[TO CONFIRM]','Wed 14 Oct'],
 ['Route (A Jal Mahal ridge / B Sumel–Jamdoli)','A needs permission; B needs recce','Fri 16 Oct'],
 ['Forest permission (DCF Wildlife Jaipur)','Not received — do not claim','Fri 16 Oct 6 pm'],
 ['Price','DECIDED: ₹399 paid test incl. water + refreshments (transport excl.)','Done'],
 ['[PARTNER] name, roles, split','TBD','Before posting'],
 ['UPI ID for collection','TBD','Before posting'],
 ['Refreshment buy + carry plan','Draft in run sheet (cost = estimate)','Sat 17 Oct'],
 ['Host / helper','Shubham hosts; helper TBD','Wed 14 Oct'],
 ['Insurance','None assumed — flag for paid event','Before 18 Oct']])
d.save('Shanivar_Shikhar_Launch_Package_18Oct2026.docx')

cols=['Name','WhatsApp number','Age','Fitness level','Medical conditions','Emergency contact','How they heard','Status','Paid (Rs 399)','UPI ref / UTR','Screenshot (Y/N)','Arrived','Feedback','Notes']
with open('shanivar_tracker_18oct.csv','w',newline='') as f: csv.writer(f).writerow(cols)
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Registrations'
ws.append(cols)
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='7A4A1E')
for i,w in enumerate([20,16,6,16,24,24,16,14,12,18,12,10,30,20]): ws.column_dimensions[chr(65+i)].width=w
for col,opts in [('H','Confirmed,Waitlist,Cancelled,No-show'),('I','399,0,Refunded'),('K','Y,N'),('L','Y,N'),('D','Easy,Moderate,Regular hiker')]:
    dv=DataValidation(type='list',formula1=f'"{opts}"'); ws.add_data_validation(dv); dv.add(f'{col}2:{col}30')
ws.freeze_panes='A2'
s=wb.create_sheet('Summary')
for r in [['Confirmed','=COUNTIF(Registrations!H:H,"Confirmed")'],['Waitlist','=COUNTIF(Registrations!H:H,"Waitlist")'],['Arrived','=COUNTIF(Registrations!L:L,"Y")'],['Paid seats','=COUNTIF(Registrations!I:I,399)'],['Revenue (Rs)','=SUMIF(Registrations!I:I,399)'],['Cost estimate (Rs, ESTIMATE)','850-1150'],['Cap',8],['Min paid by Fri 16 Oct 6pm',6]]: s.append(r)
wb.save('shanivar_tracker_18oct.xlsx')
